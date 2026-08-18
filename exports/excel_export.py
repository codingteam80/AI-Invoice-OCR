"""Exports a list of invoices to a formatted Excel workbook."""
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.drawing.image import Image as XLImage
from config.settings import settings
from exports.common import EXPORT_COLUMNS, invoice_export_row, invoice_totals_row, LINE_ITEM_COLUMNS, invoice_line_items_rows


def export_to_excel(invoices: list[dict], filename: str = "invoices_export.xlsx", chart_path: str | None = None) -> list[str]:
    rows = [invoice_export_row(inv) for inv in invoices]
    rows.append(invoice_totals_row(invoices))
    df = pd.DataFrame(rows, columns=EXPORT_COLUMNS)
    out_path = Path(settings.EXPORT_DIR) / filename

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Invoices")
        ws = writer.sheets["Invoices"]
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        # Bold the TOTAL row (last data row) so the sums stand out.
        totals_fill = PatternFill(start_color="E4E9E3", end_color="E4E9E3", fill_type="solid")
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = totals_fill

        li_rows = invoice_line_items_rows(invoices)
        if li_rows:
            li_df = pd.DataFrame(li_rows, columns=LINE_ITEM_COLUMNS)
            li_df.to_excel(writer, index=False, sheet_name="Line Items")
            li_ws = writer.sheets["Line Items"]
            for cell in li_ws[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = header_fill
            for col in li_ws.columns:
                max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                li_ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        if chart_path:
            chart_ws = writer.book.create_sheet("Chart")
            chart_ws.add_image(XLImage(chart_path), "A1")

    return [str(out_path)]
