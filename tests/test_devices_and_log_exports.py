import csv
import io
import sys
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pytest
from PIL import Image
from services import scanner_service as scan
from services.log_export_service import export_logs
from services.print_report_service import build_print_report, report_signature


def png():
    stream = io.BytesIO()
    Image.new('RGB', (80, 100), 'white').save(stream, format='PNG')
    return stream.getvalue()


def test_scan_uses_explicit_device_no_shell_and_cleans_temp(monkeypatch):
    seen = {}
    monkeypatch.setattr(scan, 'find_console', lambda: 'NAPS2.Console.exe')
    def run(args, timeout):
        seen['args'] = args
        seen['output'] = Path(args[-1])
        seen['output'].write_bytes(png())
        return ''
    monkeypatch.setattr(scan, '_run', run)
    name, data = scan.scan_page('CanoScan LiDE 300', 'twain')
    assert name.endswith('.png') and data == png()
    assert seen['args'][seen['args'].index('--device') + 1] == 'CanoScan LiDE 300'
    assert '--noprofile' in seen['args'] and '--disableocr' in seen['args']
    assert not seen['output'].exists()


def test_discovery_and_error(monkeypatch):
    monkeypatch.setattr(scan, 'find_console', lambda: 'scanner.exe')
    monkeypatch.setattr(scan, '_run', lambda *a: 'Canon LiDE 300\nCanon LiDE 300\nNetwork scanner\n')
    assert scan.list_scanners('twain') == ['Canon LiDE 300', 'Network scanner']
    with pytest.raises(scan.ScannerError): scan.list_scanners('invalid')
    with pytest.raises(scan.ScannerError): scan.scan_page('Canon', 'twain', dpi=999)
    with pytest.raises(scan.ScannerError): scan.scan_page('Canon', 'twain')  # No output produced.


def test_busy_device_and_failed_command():
    scan._SCAN_LOCK.acquire()
    try:
        with pytest.raises(scan.ScannerError, match='busy'): scan._run(['unused'], 1)
    finally: scan._SCAN_LOCK.release()
    with pytest.raises(scan.ScannerError, match='failed'):
        scan._run([sys.executable, '-c', 'import sys;sys.exit(1)'], 5)
    assert not scan._SCAN_LOCK.locked()


def test_timeout_releases_device():
    with pytest.raises(scan.ScannerError, match='timed out'):
        scan._run([sys.executable, '-c', 'import time;time.sleep(5)'], 0.1)
    assert not scan._SCAN_LOCK.locked()


def test_combine_pages():
    import fitz
    pages = [('first.png', png()), ('second.png', png())]
    assert scan.prepare_scans(pages, False) == pages
    name, data = scan.prepare_scans(pages, True)[0]
    assert name.endswith('.pdf')
    with fitz.open(stream=data, filetype='pdf') as doc: assert len(doc) == 2


def test_export_preserves_selection_and_text():
    from openpyxl import load_workbook
    rows = [dict(id=27, user='=1+1', action='EDIT', entity_id=0, details='first\nsecond, value')]
    records = list(csv.DictReader(io.StringIO(export_logs(rows, 'csv').decode('utf-8-sig'))))
    assert len(records) == 1 and records[0]['Log ID'] == '27'
    assert records[0]['User'] == "'=1+1" and records[0]['Record ID'] == '0'
    assert records[0]['Details'] == 'first\nsecond, value'
    book = load_workbook(io.BytesIO(export_logs(rows, 'xlsx')))
    assert book.active['C2'].value == '=1+1' and book.active['C2'].data_type == 's'
    assert book.active.max_row == 2


def test_print_signature_invalidation():
    invoices = [dict(id=1, total_amount=100.0)]
    before = report_signature(invoices)
    invoices[0]['total_amount'] = 101
    assert report_signature(invoices) != before
